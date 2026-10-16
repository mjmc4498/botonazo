#!/usr/bin/env python3
"""Genera presentacion/botonazo.pptx con capturas reales del producto.

Requiere python-pptx y Pillow (pip install python-pptx pillow). Las imágenes salen de
presentacion/img (capturas de la tienda, pantallas del celular en modo demostración).

    python3 scripts/presentacion.py
"""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "presentacion" / "img"
OUT = ROOT / "presentacion" / "botonazo.pptx"
GEN = IMG / "generadas"
GEN.mkdir(parents=True, exist_ok=True)

BG = RGBColor(0x0B, 0x0B, 0x0C)
CARD = RGBColor(0x18, 0x18, 0x1B)
CARD2 = RGBColor(0x22, 0x22, 0x26)
TEXT = RGBColor(0xF5, 0xF5, 0xF7)
MUTED = RGBColor(0xA1, 0xA1, 0xA6)
ACCENT = RGBColor(0xFF, 0x9F, 0x0A)
FONT = "Arial"

# ---------------------------------------------------------------- imágenes

def fondo():
    path = GEN / "fondo.png"
    w, h = 1920, 1080
    base = Image.new("RGB", (w, h), (11, 11, 12))
    glow = Image.new("L", (w, h), 0)
    d = ImageDraw.Draw(glow)
    d.ellipse((1150, 520, 2250, 1400), fill=120)
    d.ellipse((-400, -500, 600, 300), fill=40)
    glow = glow.filter(ImageFilter.GaussianBlur(220))
    orange = Image.new("RGB", (w, h), (255, 159, 10))
    base = Image.composite(orange, base, glow.point(lambda v: int(v * 0.32)))
    base.save(path)
    return path


def redondear(src, dest, radio):
    im = Image.open(src).convert("RGBA")
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=radio, fill=255)
    im.putalpha(mask)
    im.save(dest)
    return dest


def celular(nombre):
    """Pantalla del celular con bordes redondeados y marco oscuro."""
    dest = GEN / f"{nombre}-marco.png"
    pantalla = Image.open(IMG / f"{nombre}.png").convert("RGBA")
    sw, sh = pantalla.size
    marco = 28
    w, h = sw + marco * 2, sh + marco * 2
    out = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    cuerpo = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(cuerpo)
    d.rounded_rectangle((0, 0, w - 1, h - 1), radius=118, fill=(26, 26, 29, 255), outline=(70, 70, 76, 255), width=3)
    out.alpha_composite(cuerpo)
    mask = Image.new("L", (sw, sh), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, sw - 1, sh - 1), radius=92, fill=255)
    pantalla.putalpha(mask)
    out.alpha_composite(pantalla, (marco, marco))
    out.save(dest)
    return dest


# ---------------------------------------------------------------- utilidades de diapositivas

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
FONDO = fondo()


def nueva():
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(str(FONDO), 0, 0, prs.slide_width, prs.slide_height)
    return s


def texto(s, x, y, w, h, contenido, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=1.15):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    for i, p_text in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if i:
            p.space_before = Pt(size * 0.5)
        run = p.add_run()
        run.text = p_text
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        if spacing is not None:
            rPr = run._r.get_or_add_rPr()
            rPr.set("spc", str(spacing))
    return box


def cabecera(s, kicker, titulo, y=0.62, w=11.8, size=36):
    texto(s, 0.7, y, 8, 0.3, kicker.upper(), size=12, color=ACCENT, bold=True, spacing=200)
    texto(s, 0.7, y + 0.38, w, 1.2, titulo, size=size, bold=True, line=1.05)


def tarjeta(s, x, y, w, h, titulo, cuerpo, numero=None, color=CARD):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    top = y + 0.3
    if numero is not None:
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.3), Inches(top), Inches(0.42), Inches(0.42))
        c.fill.solid()
        c.fill.fore_color.rgb = ACCENT
        c.line.fill.background()
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(numero)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = RGBColor(0, 0, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        top += 0.62
    # Arial negrita de 17 pt: unos 6 caracteres por pulgada. Si el título ocupa dos
    # líneas, el texto empieza más abajo.
    lineas = 2 if len(titulo) > (w - 0.6) * 6.2 else 1
    texto(s, x + 0.3, top, w - 0.6, 0.3 * lineas + 0.2, titulo, size=17, bold=True, line=1.05)
    cuerpo_y = top + 0.18 + 0.3 * lineas
    texto(s, x + 0.3, cuerpo_y, w - 0.6, h - (cuerpo_y - y) - 0.2, cuerpo, size=13, color=MUTED, line=1.25)
    return shp


def foto(s, path, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def lista(s, x, y, w, items, size=15, gap=0.62):
    for i, (titulo, detalle) in enumerate(items):
        yy = y + i * gap
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.09), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()
        texto(s, x + 0.3, yy, w - 0.3, 0.3, titulo, size=size, bold=True)
        texto(s, x + 0.3, yy + 0.27, w - 0.3, 0.3, detalle, size=size - 3, color=MUTED)


# ---------------------------------------------------------------- imágenes preparadas

pantallas = {n: redondear(IMG / f"{n}.jpg", GEN / f"{n}-r.png", 28) for n in ["desk-message", "desk-lower", "desk-toast", "desk-confetti", "desk-clean"]}
editor = redondear(IMG / "editor-dark.jpg", GEN / "editor-dark-r.png", 30)
popup = redondear(IMG / "popup.jpg", GEN / "popup-r.png", 34)
phones = {n: celular(n) for n in ["phone-inicio", "phone-obs", "phone-atajos", "phone-tabs", "phone-viewer", "phone-shot"]}

# ---------------------------------------------------------------- 1. portada

s = nueva()
foto(s, IMG / "logo-512.png", 0.7, 1.55, w=1.25)
texto(s, 0.7, 3.0, 7, 1.1, "Botonazo", size=66, bold=True)
texto(s, 0.7, 4.15, 6.6, 1, "Tu deck de botones para Chrome, en el celular", size=24, color=MUTED, line=1.2)
texto(s, 0.7, 6.55, 7, 0.4, "Miguel Mogrovejo · Octubre de 2026", size=13, color=MUTED)
foto(s, phones["phone-inicio"], 9.15, 0.55, h=6.4)

# ---------------------------------------------------------------- 2. el reto

s = nueva()
cabecera(s, "El reto", "Transmitir y trabajar sin perder el ritmo")
for i, (t, c) in enumerate([
    ("Saltos constantes", "Cambiar de escena, de pestaña o de ventana interrumpe el directo y la concentración."),
    ("Hardware caro", "Un deck físico de botones cuesta, ocupa espacio en el escritorio y no siempre está a mano."),
    ("Atajos que se olvidan", "Cada herramienta tiene sus combinaciones; recordarlas todas en vivo es imposible."),
    ("Todo disperso", "Capturas, grabaciones, OBS, música y mensajes viven en lugares distintos."),
]):
    tarjeta(s, 0.7 + i * 3.03, 2.25, 2.83, 3.6, t, c, numero=i + 1)

# ---------------------------------------------------------------- 3. la solución

s = nueva()
cabecera(s, "La solución", "Tu celular, convertido en un deck de botones", w=6.6)
texto(s, 0.7, 2.55, 5.9, 1.2, "Una extensión de Chrome y una página web: diseñas tus botones, escaneas un código QR y tu celular controla el navegador con un toque.", size=16, color=MUTED, line=1.3)
lista(s, 0.7, 3.95, 5.9, [
    ("Sin instalar nada en el celular", "El QR abre el deck en el navegador; no hay apps ni cuentas."),
    ("Todo en un mismo lugar", "Capturas, grabación, OBS, pestañas, multimedia, mensajes y webhooks."),
    ("Directo y cifrado", "WebRTC de punto a punto entre tu celular y tu computadora."),
], gap=0.78)
foto(s, pantallas["desk-confetti"], 6.9, 1.35, w=5.8)
foto(s, phones["phone-inicio"], 6.35, 3.1, h=4.1)

# ---------------------------------------------------------------- 4. cómo se usa

s = nueva()
cabecera(s, "Cómo se usa", "Tres pasos y estás al mando")
for i, (t, c) in enumerate([
    ("Diseña tu deck", "Elige la acción de cada botón, su ícono o emoji y su color en el editor de la extensión."),
    ("Escanea el código", "Haz clic en el ícono de Botonazo (Alt+Shift+D) y apunta la cámara del celular al QR."),
    ("¡Dale al botonazo!", "Tu deck aparece en el celular y cada toque se ejecuta al instante en Chrome."),
]):
    tarjeta(s, 0.7 + i * 2.75, 2.25, 2.55, 4.1, t, c, numero=i + 1)
foto(s, popup, 9.35, 1.25, h=5.75)

# ---------------------------------------------------------------- 5. editor

s = nueva()
cabecera(s, "El editor", "Diseña tu deck a tu medida", w=4.5)
lista(s, 0.7, 2.35, 3.9, [
    ("46 acciones", "En 9 categorías, con sus opciones."),
    ("103 íconos o emojis", "Y 12 colores para cada botón."),
    ("Páginas y columnas", "Hasta 12 páginas de 40 botones."),
    ("Arrastrar y soltar", "Ordena o mueve botones entre páginas."),
    ("Macros", "Hasta 20 acciones con pausas."),
    ("Importar y exportar", "Comparte tu deck como JSON."),
], gap=0.72)
foto(s, editor, 4.75, 1.0, w=8.1)

# ---------------------------------------------------------------- 6. capturas y grabación

s = nueva()
cabecera(s, "Capturas y grabación", "Tu pantalla, en la palma de tu mano", w=6)
lista(s, 0.7, 2.35, 5.4, [
    ("Captura de pantalla", "A Descargas › Botonazo y al celular, lista para guardar o compartir."),
    ("Grabar la pestaña", "Video MP4 con sonido; el botón muestra el cronómetro."),
    ("Grabar la pantalla o una ventana", "Desde el estudio, con el selector de Chrome."),
    ("Pantalla en vivo", "Mira la pestaña en tu celular con video en tiempo real."),
], gap=0.85)
foto(s, phones["phone-viewer"], 6.75, 0.6, h=6.3)
foto(s, phones["phone-shot"], 9.85, 0.6, h=6.3)

# ---------------------------------------------------------------- 7. OBS

s = nueva()
cabecera(s, "Transmisiones", "Transmite y graba con OBS Studio", w=6.4)
lista(s, 0.7, 2.35, 5.6, [
    ("En vivo y grabación", "Inicia, detén y pausa; con cronómetro en el botón."),
    ("Escenas y fuentes", "Cambia de escena y muestra u oculta la cámara."),
    ("Audio", "Silencia el micrófono o cualquier fuente."),
    ("Repetición y cámara virtual", "Guarda el último momento o activa la cámara virtual."),
    ("Estado real", "Los botones se iluminan según lo que pasa en OBS."),
], gap=0.8)
foto(s, phones["phone-obs"], 6.45, 0.6, h=6.3)
foto(s, pantallas["desk-lower"], 9.4, 2.2, w=3.55)
texto(s, 9.4, 4.3, 3.55, 1.2, "Los carteles se dibujan sobre la página y salen en la transmisión.", size=12, color=MUTED, line=1.25)

# ---------------------------------------------------------------- 8. pestañas, ventanas y multimedia

s = nueva()
cabecera(s, "Navegador", "Pestañas, ventanas y multimedia", w=6.4)
lista(s, 0.7, 2.35, 5.8, [
    ("Lista de pestañas en el celular", "Salta a cualquiera o ciérrala con un toque."),
    ("Siguiente, anterior y la que suena", "Navega sin mirar la computadora."),
    ("Ventanas", "Cambia de ventana, abre una nueva, pantalla completa o minimiza."),
    ("Multimedia", "Play, pausa, ±10 s, volumen, silencio y velocidad."),
    ("Página", "Zoom, desplazar, atrás y adelante."),
], gap=0.8)
foto(s, phones["phone-tabs"], 6.75, 0.6, h=6.3)
foto(s, phones["phone-atajos"], 9.85, 0.6, h=6.3)

# ---------------------------------------------------------------- 9. mensajes en pantalla

s = nueva()
cabecera(s, "Mensajes en pantalla", "Habla con tu audiencia sin decir nada")
for i, (n, t) in enumerate([
    ("desk-message", "Cartel a pantalla completa"),
    ("desk-lower", "Barra inferior"),
    ("desk-toast", "Aviso rápido"),
    ("desk-confetti", "Lluvia de confeti"),
]):
    x = 0.7 + (i % 2) * 6.05
    y = 1.95 + (i // 2) * 2.7
    foto(s, pantallas[n], x, y, w=4.0)
    texto(s, x + 4.2, y + 0.85, 1.75, 0.8, t, size=15, bold=True, line=1.15)

# ---------------------------------------------------------------- 10. automatización

s = nueva()
cabecera(s, "Automatización", "Miles de funciones, las que tú quieras")
for i, (t, c) in enumerate([
    ("Atajos de teclado", "Graba cualquier combinación y envíala a la página: YouTube, Meet, Slides, editores web…"),
    ("Textos y portapapeles", "Escribe saludos o respuestas predefinidas, o copia enlaces al portapapeles."),
    ("Webhooks", "Conecta Streamer.bot, Home Assistant, Zapier, IFTTT o Discord con una petición HTTP."),
    ("Macros", "Encadena acciones: pantalla completa + reproducir, escena + mensaje, y lo que imagines."),
]):
    tarjeta(s, 0.7 + (i % 2) * 6.05, 2.2 + (i // 2) * 2.35, 5.85, 2.1, t, c)

# ---------------------------------------------------------------- 11. cómo funciona

s = nueva()
cabecera(s, "Cómo funciona", "Una conexión directa y cifrada")


def nodo(x, y, w, h, titulo, detalle, color=CARD):
    tarjeta(s, x, y, w, h, titulo, detalle, color=color)


nodo(0.7, 2.6, 3.4, 2.1, "Celular", "Deck web en el navegador. Solo envía el id del botón pulsado.")
nodo(4.95, 2.6, 3.4, 2.1, "Computadora", "Chrome con Botonazo: resuelve el botón y ejecuta la acción.")
nodo(9.2, 2.6, 3.4, 2.1, "OBS Studio", "WebSocket local para escenas, transmisión y grabación.")
for x in (4.15, 8.4):
    a = s.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(x), Inches(3.4), Inches(0.75), Inches(0.45))
    a.fill.solid()
    a.fill.fore_color.rgb = ACCENT
    a.line.fill.background()
texto(s, 0.7, 5.05, 11.9, 0.4, "WebRTC cifrado de punto a punto · señalización PeerJS solo para el enlace inicial · retransmisión de Cloudflare si la red lo bloquea", size=13, color=MUTED, align=PP_ALIGN.CENTER)
nodo(3.2, 5.55, 6.9, 1.5, "Seguro por diseño", "Las direcciones, textos y contraseñas de tus botones nunca salen de la computadora.", color=CARD2)

# ---------------------------------------------------------------- 12. privacidad

s = nueva()
cabecera(s, "Privacidad", "Tus datos se quedan contigo")
for i, (t, c) in enumerate([
    ("Sin cuentas ni analítica", "Botonazo no recopila datos personales, no muestra anuncios y no usa cookies."),
    ("Código QR con clave", "La clave viaja en el fragmento # y se renueva en cada sesión."),
    ("Archivos en tus equipos", "Capturas y videos se guardan en tu computadora o tu celular, nunca en un servidor."),
    ("Permisos a pedido", "Los scripts entran a la página solo cuando pulsas un botón que los necesita."),
]):
    tarjeta(s, 0.7 + i * 3.03, 2.25, 2.83, 3.6, t, c, numero=i + 1)

# ---------------------------------------------------------------- 13. estado y publicación

s = nueva()
cabecera(s, "Estado", "Listo para la Chrome Web Store")
for i, (t, c) in enumerate([
    ("64 comprobaciones", "Prueba de punta a punta con la extensión, el celular, OBS simulado, grabación y video en vivo reales."),
    ("Paquete listo", "dist/botonazo-1.0.0.zip validado: manifiesto, íconos, idioma y scripts."),
    ("Ficha completa", "Textos, 5 capturas, mosaicos, privacidad y justificación de permisos."),
]):
    tarjeta(s, 0.7 + i * 4.05, 2.25, 3.85, 2.6, t, c, numero=i + 1)
texto(s, 0.7, 5.35, 11.9, 1.2, ["Siguiente paso: publicar el repositorio con GitHub Pages (mjmc4498.github.io/botonazo) y subir el ZIP en el panel del desarrollador."], size=16, color=MUTED, line=1.3)

# ---------------------------------------------------------------- 14. cierre

s = nueva()
foto(s, IMG / "logo-512.png", 0.7, 1.7, w=1.25)
texto(s, 0.7, 3.15, 8, 1.1, "¡Dale al botonazo!", size=56, bold=True)
texto(s, 0.7, 4.3, 7.4, 1.4, ["Diseña tu deck, escanea el QR y controla Chrome desde tu celular.", "mjmc4498.github.io/botonazo"], size=20, color=MUTED, line=1.25)
texto(s, 0.7, 6.55, 7, 0.4, "Creado por Miguel Mogrovejo", size=13, color=MUTED)
foto(s, phones["phone-obs"], 9.15, 0.55, h=6.4)

prs.save(OUT)
print("  ", OUT.relative_to(ROOT), f"({len(prs.slides)} diapositivas)")
